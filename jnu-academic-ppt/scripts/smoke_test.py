# -*- coding: utf-8 -*-
"""跨主题冒烟测试：构建、打开、静态体检，并验证缺图默认失败。"""
from __future__ import annotations

import json
import tempfile
import zipfile
from pathlib import Path

from pptx import Presentation
from pptx.oxml.ns import qn

from build_deck import ROOT, build
from check_deck import check


THEMES = ["jnu-teal", "jnu-defense", "jnu-rigor", "jnu-crisp", "jnu-red"]
EXAMPLES = [ROOT / "examples" / "组会示例.deck.json",
            ROOT / "examples" / "答辩示例.deck.json"]


TEMPLATE_AUTHORS = ("wz z", "宏馨", "水 王", "堇")


def _probe(out_dir, name, meta=None, slides=None, qa=None):
    deck = {"meta": {"title": "回归测试", "presenter": "测试同学", **(meta or {})},
            "slides": slides or [{"layout": "cover"}, {"layout": "closing"}]}
    if qa is not None:
        deck["qa"] = qa
    p = out_dir / f"{name}.deck.json"
    p.write_text(json.dumps(deck, ensure_ascii=False), encoding="utf-8")
    built, warnings = build(p, out_dir / f"{name}.pptx", "jnu-teal")
    return Presentation(str(built)), warnings


def _notes(slide):
    return slide.notes_slide.notes_text_frame.text if slide.has_notes_slide else ""


def _regressions(out_dir):
    """针对已修过的 bug 的回归检查。"""
    fails = []

    # 1. 普通页讲稿备注不能被截断
    long_note = "讲稿" * 600
    prs, _ = _probe(out_dir, "notes", slides=[
        {"layout": "cover"},
        {"layout": "bullets", "title": "备注", "bullets": ["要点"], "notes": long_note},
        {"layout": "closing"}])
    if _notes(prs.slides[1]) != long_note:
        fails.append("普通页备注被截断")
    else:
        print("OK  普通页备注完整保留")

    # 2. 问答备份：按整题取舍、不超预算、报告丢弃题数
    qa = [{"question": f"第{i}题：为什么这样设计实验？", "answer": "回答" * 60}
          for i in range(1, 13)]
    prs, warns = _probe(out_dir, "qa", qa=qa)
    text = _notes(prs.slides[-1])
    body = text.split("（另有")[0]
    if len(body) > 600 + 40 or "另有" not in text or not any("问答备份" in w for w in warns):
        fails.append(f"问答备份预算异常（{len(text)} 字）")
    elif text.count("Q：") + int(text.split("另有 ")[1].split(" 题")[0]) != 12:
        fails.append("问答备份题数对不上")
    else:
        print(f"OK  问答备份 {text.count('Q：')} 题入备注，其余注明见对话")

    # 3. 没写 presentation_mode 时按时长自动紧凑
    sec = [{"layout": "cover"}, {"layout": "toc"}, {"layout": "section", "title": "一"},
           {"layout": "bullets", "title": "内容", "bullets": ["x"]}, {"layout": "closing"}]
    short, _ = _probe(out_dir, "short", meta={"duration_minutes": 8}, slides=sec)
    long_, _ = _probe(out_dir, "long", meta={"duration_minutes": 20}, slides=sec)
    forced, _ = _probe(out_dir, "forced", slides=sec,
                       meta={"duration_minutes": 8, "presentation_mode": "standard"})
    got = (len(short.slides), len(long_.slides), len(forced.slides))
    if got != (3, 5, 5):
        fails.append(f"紧凑模式判断异常：8 分钟/20 分钟/显式 standard = {got} 页")
    else:
        print("OK  10 分钟以内自动紧凑，显式设置优先")

    # 4. 成品文件属性不能带模板原作者
    cp = short.core_properties
    leaked = [a for a in TEMPLATE_AUTHORS
              if a in (cp.author or "") or a in (cp.last_modified_by or "")]
    if leaked or cp.author != "测试同学":
        fails.append(f"文件属性作者异常：{cp.author!r} / {cp.last_modified_by!r}")
    else:
        print("OK  文件属性作者按 meta 重写")
    for stencil in (ROOT / "themes").glob("*/stencil.pptx"):
        with zipfile.ZipFile(stencil) as z:
            core = z.read("docProps/core.xml").decode("utf-8")
        if any(a in core for a in TEMPLATE_AUTHORS):
            fails.append(f"{stencil.parent.name}/stencil.pptx 仍带原作者姓名")

    # 5. 图表跟随主题、表格单元格垂直居中
    prs, _ = _probe(out_dir, "chart", slides=[
        {"layout": "cover"},
        {"layout": "chart", "title": "图表", "lead": "结论句",
         "chart": {"categories": ["a", "b"], "series": [{"name": "s", "values": [1, 2]}]}},
        {"layout": "table", "title": "表格", "header": ["h1", "h2"], "rows": [["1", "2"]]},
        {"layout": "closing"}])
    chart = next(sh.chart for sh in prs.slides[1].shapes if sh.has_chart)
    has_lead = any(sh.has_text_frame and sh.text_frame.text == "结论句"
                   for sh in prs.slides[1].shapes)
    table = next(sh.table for sh in prs.slides[2].shapes if sh.has_table)
    anchor = table.cell(1, 0)._tc.tcPr.get("anchor")
    if chart.plots[0].vary_by_categories or not has_lead or anchor != "ctr":
        fails.append("图表/表格样式回退：彩虹色、lead 未渲染或单元格未垂直居中")
    else:
        print("OK  图表跟随主题配色并渲染 lead，表格单元格垂直居中")

    # 5b. 横向条形图：按书写顺序从上到下、负值不反色、正负语义色
    prs, _ = _probe(out_dir, "bar", slides=[
        {"layout": "cover"},
        {"layout": "chart", "title": "变化量",
         "chart": {"type": "bar", "categories": ["先", "后"],
                   "series": [{"name": "d", "values": [-1.5, 0.5]}],
                   "number_format": "+0.0;-0.0;0.0", "data_labels": True,
                   "sign_colors": True}},
        {"layout": "closing"}])
    chart = next(sh.chart for sh in prs.slides[1].shapes if sh.has_chart)
    ser = chart.plots[0].series[0]
    dpts = ser._element.findall(qn("c:dPt"))
    ok_order = list(chart.plots[0].categories) == ["后", "先"]   # PowerPoint 自下而上画
    ok_invert = dpts and all(d.find(qn("c:invertIfNegative")) is not None
                             and d.find(qn("c:invertIfNegative")).get("val") == "0"
                             for d in dpts)
    if not (ok_order and ok_invert and chart.plots[0].has_data_labels):
        fails.append("条形图回退：顺序、负值反色或数值标签异常")
    else:
        print("OK  条形图按书写顺序排列，负值不反色，带数值标签")

    # 6. Windows PowerShell 5.1 按 ANSI 读无 BOM 的脚本，中文会把语法搞坏
    for ps1 in (ROOT / "scripts").glob("*.ps1"):
        if not ps1.read_bytes().startswith(b"\xef\xbb\xbf"):
            fails.append(f"{ps1.name} 缺少 UTF-8 BOM，Windows PowerShell 5.1 会解析失败")
    return fails


def main():
    failures = []
    with tempfile.TemporaryDirectory(prefix="jnu-ppt-smoke-") as tmp:
        out_dir = Path(tmp)
        for example in EXAMPLES:
            for theme in THEMES:
                out = out_dir / f"{example.stem}-{theme}.pptx"
                try:
                    built, warnings = build(example, out, theme)
                    prs = Presentation(str(built))
                    if len(prs.slides) < 3:
                        failures.append(f"{example.name}/{theme}: 页数异常")
                    report = check(built, theme, deep=False)
                    if report.errors:
                        failures.append(
                            f"{example.name}/{theme}: {len(report.errors)} 个静态 ERROR")
                    print(f"OK  {example.name:<18} {theme:<12} "
                          f"{len(prs.slides):>2} 页  {len(warnings)} 警告")
                except Exception as exc:
                    failures.append(f"{example.name}/{theme}: {exc}")

        probe = json.loads(EXAMPLES[0].read_text(encoding="utf-8"))
        probe["slides"] = [
            {"layout": "cover"},
            {"layout": "image-full", "title": "缺图测试",
             "image": "__definitely_missing__.png"},
        ]
        probe_path = out_dir / "missing.deck.json"
        probe_path.write_text(json.dumps(probe, ensure_ascii=False), encoding="utf-8")
        try:
            build(probe_path, out_dir / "should-not-exist.pptx", "jnu-teal")
            failures.append("缺图测试：默认没有阻止导出")
        except SystemExit:
            print("OK  缺图默认阻止导出")

        failures += _regressions(out_dir)

    if failures:
        print("\nFAILED")
        for item in failures:
            print(" - " + item)
        raise SystemExit(1)
    print("\n全部冒烟测试通过。")


if __name__ == "__main__":
    main()
