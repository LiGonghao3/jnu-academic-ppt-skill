# -*- coding: utf-8 -*-
"""解剖一套 .pptx 模板，给「收编成新主题包」提供依据。

会输出：
  · 画布尺寸、母版/版式清单
  · 每页的形状树（可穿透组合），含名称、坐标、文字预览
  · 实际用到的颜色频次（注意：theme1.xml 里往往是 Office 默认配色，
    真正的品牌色藏在各页形状里，所以这里统计的是页面实际用色）
  · 实际用到的字体频次，以及有没有内嵌字体（内嵌字体动辄几十 MB）
  · 媒体清单和有效分辨率

用法：
    python inspect_template.py 模板.pptx [--all] [--slides 1,2,5] [--colors] [--media]
"""
from __future__ import annotations

import argparse
import collections
import posixpath
import re
import zipfile
from pathlib import Path

from lxml import etree
from pptx import Presentation

NS = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
EMU = 914400


def head(txt, ch="="):
    print("\n" + ch * 78)
    print(txt)
    print(ch * 78)


def deck_overview(prs):
    print("画布：%.2f × %.2f 英寸  (%d × %d EMU)" % (
        prs.slide_width / EMU, prs.slide_height / EMU, prs.slide_width, prs.slide_height))
    print("母版 %d 套，幻灯片 %d 页" % (len(prs.slide_masters), len(prs.slides)))
    for mi, m in enumerate(prs.slide_masters):
        print("  母版 %d 的版式（%d 个）：%s" % (
            mi, len(m.slide_layouts), "、".join(l.name for l in m.slide_layouts)))


def walk(shapes, depth=0):
    for sh in shapes:
        try:
            pos = "(%6.2f,%6.2f) %6.2f×%-6.2f" % (
                sh.left / EMU, sh.top / EMU, sh.width / EMU, sh.height / EMU)
        except Exception:
            pos = " " * 30
        txt = ""
        if sh.has_text_frame:
            txt = " | ".join(p.text for p in sh.text_frame.paragraphs if p.text)[:56]
        print("   " + "  " * depth + "%-13s %-24s %s  %s" % (
            sh.shape_type, sh.name[:24], pos, txt))
        if sh.shape_type == 6:  # GROUP：子形状坐标在组内坐标系，仅供参考
            walk(sh.shapes, depth + 1)


def slides(prs, want=None):
    for i, s in enumerate(prs.slides, 1):
        if want and i not in want:
            continue
        head("第 %d 页   版式=%s" % (i, s.slide_layout.name), "-")
        walk(s.shapes)


def colors_fonts(path):
    cnt, fonts, sizes = collections.Counter(), collections.Counter(), collections.Counter()
    with zipfile.ZipFile(path) as z:
        for n in z.namelist():
            if not n.endswith(".xml"):
                continue
            if not (n.startswith("ppt/slides/slide") or n.startswith("ppt/slideLayouts/")
                    or n.startswith("ppt/slideMasters/")):
                continue
            x = z.read(n).decode("utf8", "ignore")
            for m in re.finditer(r'srgbClr val="([0-9A-Fa-f]{6})"', x):
                cnt["#" + m.group(1).upper()] += 1
            for m in re.finditer(r'typeface="([^"]+)"', x):
                if not m.group(1).startswith("+"):
                    fonts[m.group(1)] += 1
            for m in re.finditer(r'\ssz="(\d+)"', x):
                sizes[int(m.group(1)) // 100] += 1
        embedded = [(n, z.getinfo(n).file_size) for n in z.namelist()
                    if n.startswith("ppt/fonts/")]

    head("页面实际用色（频次越高越可能是主色）")
    for c, k in cnt.most_common(16):
        print("   %s  ×%d" % (c, k))
    head("实际用到的字体")
    for f, k in fonts.most_common(12):
        print("   %-28s ×%d" % (f, k))
    head("字号分布（pt）")
    print("   " + "、".join("%d(×%d)" % (s, k) for s, k in sizes.most_common(12)))
    if embedded:
        total = sum(s for _, s in embedded) / 1048576
        print("\n⚠ 内嵌了 %d 个字体文件，共 %.1f MB。" % (len(embedded), total))
        print("  slim_pptx.py 默认会剥掉它们，theme.json 里要写好回退链。")


def media(path):
    head("媒体文件")
    with zipfile.ZipFile(path) as z:
        items = sorted(((z.getinfo(n).file_size, n) for n in z.namelist()
                        if n.startswith("ppt/media/")), reverse=True)
        for size, n in items[:20]:
            dim = ""
            ext = posixpath.splitext(n)[1].lower()
            if ext in (".png", ".jpg", ".jpeg", ".bmp"):
                try:
                    import io
                    from PIL import Image
                    w, h = Image.open(io.BytesIO(z.read(n))).size
                    dim = "  %d×%d" % (w, h)
                except Exception:
                    pass
            print("   %8.1f KB  %-28s%s" % (size / 1024, n.split("/")[-1], dim))


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("pptx")
    ap.add_argument("--all", action="store_true", help="全套：结构 + 逐页 + 颜色字体 + 媒体")
    ap.add_argument("--slides", help="只看某几页，如 1,2,5")
    ap.add_argument("--colors", action="store_true")
    ap.add_argument("--media", action="store_true")
    a = ap.parse_args()

    p = Path(a.pptx)
    head(p.name)
    prs = Presentation(str(p))
    deck_overview(prs)

    want = {int(x) for x in a.slides.split(",")} if a.slides else None
    if a.all or a.slides:
        slides(prs, want)
    if a.all or a.colors:
        colors_fonts(p)
    if a.all or a.media:
        media(p)
    if not (a.all or a.slides or a.colors or a.media):
        print("\n（加 --all 看全套，或 --slides 1,2 / --colors / --media）")


if __name__ == "__main__":
    main()
